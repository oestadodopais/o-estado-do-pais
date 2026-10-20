"""Conferência de partida sobre os corpos alojados e o livro selado.

Uso: python3 conferir-partida.py <motor>
As contas independentes são comparadas com as linhas seladas antes da exportação.
"""
import hashlib
import io
import json
import re
import subprocess
import sys
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

sys.path.insert(0, sys.argv[1])
from publisher import oe2027_analise as A
import openpyxl

OUT = Path(__file__).resolve().parent
PDF = 'content/15 Orcamento do Estado 2027/source/v2a/oe-relatorio.pdf'
book = A.load(A.STUDY / 'ledger.json')
regenerated = A.build()
assert book == regenerated, 'O livro difere da regeneração dos corpos'
by = {r['id']: r for r in book['claims']}
specs = {r['id']: r for r in A.load(A.STUDY / 'extracoes.json')['rows']}
rows = []
failures = []

def numeric(value):
    value = str(value).replace(' ', '').replace('\u00a0', '').replace('−', '-')
    if value.startswith('('): value = '-' + value[1:-1]
    return Decimal(value.replace(',', '.'))

def record(id, value, expected=None, **proof):
    if expected is not None and numeric(value) != numeric(expected):
        failures.append(dict(id=id, esperado=expected, lido=value))
    entry = dict(id=id, value=value, expected=expected, **proof)
    rows.append(entry)
    return entry

def old(id, expected=None):
    r = by[id]
    entry = record(id, r['value'], expected, period=r['period'],
                   estado='selada na cabeça de partida', document=r.get('document'),
                   excerpt=r['excerpt'], derivation=r.get('derivation'), note=r.get('note'))
    d = r.get('document')
    if d:
        body, receipt = A.source(d['file'])
        assert hashlib.sha256(body).hexdigest() == d['source_sha256']
        if d['file'] == PDF:
            page = A.pages(PDF)[d['page']-1]
            s = specs.get(id)
            if s and s['kind'] == 'pdf_cell':
                entry['linha_pdf'] = [l for l in page.splitlines() if l.strip().startswith(s['label'])]
                entry['column'] = s['column']
                entry['column_label'] = s.get('column_label')
                if s.get('metric') in ('efetiva','sem-prr') and s.get('stage') in ('estimativa','orcamento'):
                    expected_column=(0 if str(r['period'])=='2026' else 2)+(1 if s['metric']=='sem-prr' else 0)
                    assert s['column']==expected_column, 'Coluna e ano divergentes: '+id
                    entry['coluna_ano_conferidos']=True
                elif id.startswith('ppp-oe2027-liquidos-'):
                    year=int(r['period'])
                    expected_column=year-(2026 if year<2040 else 2040)
                    assert s['column']==expected_column, 'Coluna e ano divergentes: '+id
                    entry['coluna_ano_conferidos']=True
            # Guardar a cabeça real, incluindo os cabeçalhos das colunas.
            entry['cabecalho_pdf'] = page.splitlines()[:16]
            if not s:
                labels = ['Despesa total consolidada', 'Despesa efetiva consolidada'] if 'saude-' in id else ['Rodoviário', 'encargos brutos', 'receitas', 'Total']
                entry['linha_pdf'] = [l for l in page.splitlines() if any(l.strip().startswith(x) for x in labels)]
    return entry

health = 'oe-2027-saude-despesa-'
for metric, vals in [('total-consolidada', ['18125,9','17858,4','-1,5']),
                     ('efetiva-consolidada', ['18062,4','17803,1','-1,4']),
                     ('efetiva-consolidada-sem-prr', ['17291,3','17803,1','3,0'])]:
    for suffix, value in zip(['2026','2027','variacao'], vals): old(health+metric+'-'+suffix, value)
for id, expected in {
    'der-saude-prr-2026':'771,1', 'der-saude-variacao-sem-prr':'511,8',
    'der-saude-variacao-efetiva':'-259,3', 'der-saude-razao-compensacao-prr':'0,6637',
    'der-saude-reforco-total-2026':'825,7', 'der-financas-prr-2026':'1472,7',
    'der-economia-prr-2026':'1357,1', 'der-infraestruturas-e-habitacao-prr-2026':'1133,0',
    'der-ambiente-e-energia-prr-2026':'823,2', 'der-ppp-quota-rodoviario':'61,82',
    'utao-saude-efetiva-2026-orcamento':'17147', 'utao-saude-efetiva-2026-estimativa':'18062',
}.items(): old(id, expected)
for year in range(2020,2027):
    old(f'saude-total-{year}-orcamento-oe{year}', {2020:'11282,2',2025:'16853,5',2026:'17300,2'}.get(year))
    if year < 2026: old(f'saude-total-{year}-execucao-cge{year}', {2020:'11731,9',2025:'17241,5'}.get(year))
for id, expected in [('rodoviario-encargos-brutos','1152'),('rodoviario-receitas','263'),
                     ('rodoviario-encargos-liquidos','889'),('total-encargos-liquidos','1438')]:
    old('oe-2027-ppp-'+id+'-2027', expected)
for year, expected in [(2021,'1250'),(2022,'1192'),(2023,'1105'),(2024,'1012'),(2025,'843'),(2026,'942')]:
    old(f'ppp-oe{year+1}-liquidos-{year}', expected)
for year, expected in [(2028,'705'),(2029,'648'),(2030,'-220'),(2035,'77'),(2040,'17')]:
    old(f'ppp-oe2027-liquidos-{year}', expected)
for unit, hf, pt, eu in [('pc_gdp','tot_hf','10.20','10.16'),('eur_hab','tot_hf','2764.83','4074.88'),('pc_gdp','hf1','6.31','8.17')]:
    for geo, expected in [('pt',pt),('eu27_2020',eu)]: old(f'eurostat-hlth-sha11-hf-a-{unit}-{hf}-{geo}-2024', expected)
for id in ['afirmacao-leitao-amaro-saude-oe2027-2026-10-09','afirmacao-rtp-saude-queda-oe2027-2026-10-09',
           'afirmacao-pcp-paula-investimento-saude-2026-10-12','utao-saude-sem-prr-2027','utao-risco-execucao-sns-2027',
           'afirmacao-pcp-ppp-rodoviarias-oe2027-2026-10-09','afirmacao-pcp-paula-ppp-2026-10-12',
           'afirmacao-pcp-ppp-2026-10-15','utao-ppp-contingencias-natureza-2026']: old(id)

def cell(id, label, column, expected, page=303, occurrence=0, header=None, period=None):
    lines = A.pages(PDF)[page-1].splitlines()
    matches = [l for l in lines if l.strip().startswith(label) and re.match(r'\s{2,}', l.strip()[len(label):])]
    line = matches[occurrence]
    fields = re.split(r'\s{2,}', line.strip()[len(label):].strip())
    value = fields[column]
    d = dict(file=PDF, source_sha256=A.source(PDF)[1]['sha256'], page=page, printed_page=page-16)
    e = record(id, value, expected, period=period or ('2026' if id.endswith('2026') else '2027'),
               estado='candidata conferida, ainda não selada', document=d, linha_pdf=line,
               cabecalho_pdf=lines[:10] if header is None else header,
               column=column, occurrence=occurrence)
    by[id] = e

for slug, label, values in [
    ('pessoal','Despesas com o pessoal',['7783,1','8184,3','5,2']),
    ('bens-servicos','Aquisição de bens e serviços',['8832,9','8436,3','-4,5']),
    ('investimento','Investimento',['1022,5','678,3','-33,7']),
    ('outras-correntes','Outras despesas correntes',['32,4','303,7']),
    ('transferencias-capital','Transferências de capital',['272,9','0,0'])]:
    for suffix, column, expected in zip(['2026','2027','variacao'],[0,2,4],values):
        cell('oe-2027-saude-'+slug+'-'+suffix,label,column,expected)
cell('oe-2027-saude-investimento-variacao-sem-prr','Investimento',5,'22,2')
memory = A.pages(PDF)[302].split('por memória:')[1].splitlines()[:8]
for slug, label, values in [('impostos','Receitas de Impostos',['15116,4','15452,2']),
    ('receitas-proprias','Receitas Próprias',['1289,5','2017,6']),('fundos-europeus','Fundos Europeus',['655,0','174,6'])]:
    for column, expected in enumerate(values):cell(f'oe-2027-saude-fonte-{slug}-{2026+column}',label,column,expected,header=memory)
for suffix, column, expected in [('2026',0,'14935,8'),('2027',1,'15539,2'),('variacao',2,'4,0')]:
    cell('oe-2027-saude-transferencias-sns-'+suffix,'TRANSFERÊNCIAS SERVIÇO NACIONAL DE SAÚDE',column,expected,page=304)
cell('ppp-oe2027-ferroviario-2027','Ferroviário',1,'312',page=123)
cell('ppp-oe2027-saude-2027','Saúde',1,'254',page=123)
for id, value, page, needle in [
    ('liquidos-2027-previsao-oe2026','774',125,'Enquanto este último'),
    ('riscos-pedidos-rodoviarios','2388,2',369,'ascende a cerca de'),
    ('riscos-acrescimo','1140',369,'um acréscimo de cerca'),
    ('riscos-brisa','1122,5',369,'montante de cerca de'),
    ('riscos-rotas-algarve-principal','445',370,'60 O litígio'),
    ('riscos-rotas-algarve-mensal','1,3',370,'acrescido de um valor mensal')]:
    lines = A.pages(PDF)[page-1].splitlines()
    i = next(i for i,l in enumerate(lines) if needle in l)
    excerpt = '\n'.join(lines[max(0,i-2):i+5])
    assert value in excerpt
    by['ppp-oe2027-'+id] = record('ppp-oe2027-'+id,value,value,period='2027',
        estado='candidata conferida, ainda não selada',excerpt=excerpt,
        document=dict(file=PDF,source_sha256=A.source(PDF)[1]['sha256'],page=page,printed_page=page-16))

xfile = 'content/20 Orcamento do Estado/source/eo-anexo-agosto.xlsx'
xbody,xreceipt = A.source(xfile)
workbook = openpyxl.load_workbook(io.BytesIO(xbody),data_only=True)
for suffix, sheet, coord, expected, period in [
    ('saude-efetiva-jan-ago-2025','24 - Desp. Efetiva por PO','C26','9861,9','2025'),
    ('saude-efetiva-jan-ago-2026','24 - Desp. Efetiva por PO','D26','11822,6','2026'),
    ('saude-efetiva-variacao-homologa','24 - Desp. Efetiva por PO','F26','19,9','2026'),
    ('sns-saldo-jan-ago-2026','1 - Saldo Global Rec Desp','D16','-918,5','2026'),
    ('sns-despesa-variacao-homologa','1 - Saldo Global Rec Desp','J16','13,2','2026')]:
    s=workbook[sheet];c=s[coord]
    value=str(Decimal(str(c.value)).quantize(Decimal('.1'),rounding=ROUND_HALF_UP)).replace('.',',')
    by['dgo-2026-08-'+suffix] = record('dgo-2026-08-'+suffix,value,expected,period=period,
        estado='candidata conferida, ainda não selada',raw_value=c.value,cell=coord,sheet=sheet,number_format=c.number_format,
        label=s.cell(c.row,2).value,headers=[{x.coordinate:x.value for x in row if x.value is not None} for row in s.iter_rows(min_row=6,max_row=13)],
        document=dict(file=xfile,source_sha256=xreceipt['sha256']))

def derive(id, components, operation, expected=None, places=1):
    vals=[numeric(by[x]['value']) for x in components]
    if operation=='subtract': value=vals[0]-vals[1]
    elif operation=='change_pct':value=(vals[0]/vals[1]-1)*100
    elif operation=='sum':value=sum(vals)
    rounded=value.quantize(Decimal(1).scaleb(-places),rounding=ROUND_HALF_UP)
    by[id] = record(id,str(rounded).replace('.',','),expected,period='2027',estado='derivada candidata, ainda não selada',
        derivation=dict(components=components,values=[str(v) for v in vals],operation=operation,unrounded=str(value),places=places))

for slug, expected in [('pessoal','401,2'),('bens-servicos','-396,6'),('investimento','-344,2')]:
    derive('der-saude-'+slug+'-variacao-abs',['oe-2027-saude-'+slug+'-'+y for y in ['2027','2026']],'subtract',expected)
derive('der-saude-variacao-total-abs',[health+'total-consolidada-'+y for y in ['2027','2026']],'subtract','-267,5')
derive('der-saude-receitas-proprias-variacao-pct',['oe-2027-saude-fonte-receitas-proprias-'+y for y in ['2027','2026']],'change_pct','56,5')
derive('der-saude-2027-vs-execucao-2025-pct',[health+'total-consolidada-2027','saude-total-2025-execucao-cge2025'],'change_pct','3,6')
derive('der-ppp-brutos-todos-2027',['oe-2027-ppp-rodoviario-encargos-brutos-2027','ppp-oe2027-ferroviario-2027','ppp-oe2027-saude-2027'],'sum','1718',0)
derive('der-ppp-rodoviario-2027-vs-2021-pct',['oe-2027-ppp-rodoviario-encargos-liquidos-2027','ppp-oe2022-liquidos-2021'],'change_pct','-28,9')
for slug, expected in [('brutos','548'),('liquidos','811')]:
    derive('der-ppp-pcp-vs-'+slug,['afirmacao-pcp-ppp-rodoviarias-oe2027-2026-10-09','oe-2027-ppp-rodoviario-encargos-'+slug+'-2027'],'subtract',expected,0)
derive('der-utao-saude-reforco-2026',['utao-saude-efetiva-2026-estimativa','utao-saude-efetiva-2026-orcamento'],'subtract','915',0)
for program in ['saude','educacao','trabalho-solidariedade-e-seguranca-social','infraestruturas-e-habitacao','justica','seguranca-interna']:
    prefix=health+'efetiva-consolidada-sem-prr-' if program=='saude' else 'oe-2027-'+program+'-sem-prr-'
    if program!='saude':
        for y in ['2026','2027']:old(prefix+y)
        old('der-'+program+'-variacao-sem-prr')
    derive('der-'+program+'-variacao-sem-prr-pct',[prefix+y for y in ['2027','2026']],'change_pct','3,0' if program=='saude' else None)
# As componentes das derivadas antigas também levam a linha e o cabeçalho do PDF.
seen={r['id'] for r in rows}
for row in rows:
    for component in (row.get('derivation') or {}).get('components',[]):
        if component not in seen:
            old(component)
            seen.add(component)
checks = dict(
    efetiva_2027_igual_sem_prr=by[health+'efetiva-consolidada-2027']['value']==by[health+'efetiva-consolidada-sem-prr-2027']['value'],
    ppp_brutos_soma=by['der-ppp-brutos-todos-2027']['value'],
    execucao_superior_orcamento_por_ano={str(y):numeric(by[f'saude-total-{y}-execucao-cge{y}']['value'])>numeric(by[f'saude-total-{y}-orcamento-oe{y}']['value']) for y in range(2020,2026)},
)
sealed = {r['id']: r for r in book['claims']}
for id in ['der-saude-bens-servicos-reducao-abs', 'der-saude-investimento-reducao-abs', 'der-saude-total-reducao-abs']:
    old(id)
# A correção acrescenta provas e componentes; nenhuma linha atravessada fica fora da conferência.
seen={r['id'] for r in rows}
for id in A.load(A.ROOT/'publisher/recortes/v2d/provas.json')['rows']:
    if id not in seen:
        old(id)
        seen.add(id)
for entry in rows:
    row = sealed[entry['id']]
    if entry['value'] is not None and not row.get('value_not_numeric'):
        assert numeric(entry['value']) == numeric(row['value']), 'Livro e corpo divergem: '+entry['id']
    entry['value'] = row['value']
    entry['period'] = row['period']
    entry['estado'] = 'selada e conferida contra o corpo ou as componentes'
    entry['excerpt'] = row['excerpt']
    entry['document'] = row.get('document')
    if row.get('derivation'):
        entry['derivation_sealed'] = row['derivation']
    selector = specs.get(entry['id'])
    if selector:
        entry['selector'] = selector
        if selector.get('column_headers'):
            entry['coluna_ano_conferidos'] = True
result=dict(estado='linhas seladas; valores conferidos contra os corpos e contas refeitas a partir das componentes',
    motor_head=subprocess.check_output(['git','rev-parse','HEAD'],cwd=A.ROOT,text=True).strip(),
    site_head=subprocess.check_output(['git','rev-parse','HEAD'],cwd=OUT,text=True).strip(),
    regeneracao_integral_igual_ao_livro=True,rows=rows,checks=checks,divergencias_numericas=failures)
(OUT/'partida.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(dict(linhas=len(rows),divergencias=failures,checks=checks),ensure_ascii=False))
sys.exit(bool(failures))
